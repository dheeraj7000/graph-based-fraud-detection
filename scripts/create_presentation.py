from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RESULTS_PATH = ROOT / "results" / "metrics" / "main_results.csv"
METADATA_PATH = ROOT / "data" / "processed" / "sampled_graph_metadata.json"
OUTPUT_PATH = ROOT / "presentations" / "Graph_Based_Fraud_Detection.pptx"

NAVY = RGBColor(17, 31, 54)
BLUE = RGBColor(48, 113, 190)
TEAL = RGBColor(35, 166, 154)
GOLD = RGBColor(239, 177, 69)
INK = RGBColor(36, 48, 65)
MUTED = RGBColor(102, 116, 135)
PALE = RGBColor(239, 244, 249)
WHITE = RGBColor(255, 255, 255)
PALETTE = (BLUE, TEAL, GOLD)
FEATURES = ("tabular_sampled", "sampled_graph", "sampled_temporal")
FEATURE_LABELS = ("Sampled tabular", "Sampled graph", "Graph + 24h")
MODELS = ("Logistic", "Random Forest", "XGBoost", "MLP")


def _set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_text(
    slide,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float = 18,
    color: RGBColor = INK,
    bold: bool = False,
    font: str = "Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def _add_bullets(slide, items: list[str], x: float, y: float, width: float, height: float,
                 *, size: float = 18, color: RGBColor = INK, spacing: float = 12) -> None:
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.04)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(spacing)


def _add_footer(slide, page: int, dark: bool = False) -> None:
    color = RGBColor(183, 197, 214) if dark else MUTED
    _add_text(
        slide,
        "GRAPH-BASED FRAUD DETECTION  •  PAYSIM",
        0.55,
        7.12,
        10.9,
        0.2,
        size=8,
        color=color,
    )
    _add_text(
        slide,
        f"{page:02d}",
        12.15,
        7.08,
        0.55,
        0.25,
        size=9,
        color=color,
        align=PP_ALIGN.RIGHT,
    )


def _new_slide(prs: Presentation, title: str, page: int, kicker: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, WHITE)
    if kicker:
        _add_text(slide, kicker.upper(), 0.6, 0.36, 11.8, 0.28, size=10, color=TEAL, bold=True)
    _add_text(slide, title, 0.6, 0.72, 12.0, 0.62, size=27, color=NAVY, bold=True)
    _add_footer(slide, page)
    return slide


def _add_card(slide, title: str, body: str, x: float, y: float, width: float,
              height: float, accent: RGBColor = BLUE, body_size: float = 16) -> None:
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = PALE
    card.line.fill.background()
    _add_text(slide, title, x + 0.22, y + 0.18, width - 0.44, 0.42,
              size=17, color=accent, bold=True)
    _add_text(slide, body, x + 0.22, y + 0.7, width - 0.44, height - 0.82,
              size=body_size, color=INK)


def _validate_results(results: pd.DataFrame) -> None:
    required = {"model", "features", "pr_auc", "precision", "recall", "f1", "roc_auc"}
    missing_columns = required.difference(results.columns)
    if missing_columns:
        raise ValueError(f"Experiment results are missing columns: {sorted(missing_columns)}")
    matrix = results[results["features"].isin(FEATURES)]
    pairs = set(zip(matrix["model"], matrix["features"]))
    expected = {(model, feature) for model in MODELS for feature in FEATURES}
    if pairs != expected:
        missing = sorted(expected - pairs)
        extra = sorted(pairs - expected)
        raise ValueError(
            "The presentation requires exactly the completed 12-cell main matrix. "
            f"Missing: {missing}; unexpected: {extra}"
        )


def _add_chart_slide(prs: Presentation, results: pd.DataFrame, page: int) -> None:
    slide = _new_slide(prs, "Measured test PR-AUC", page, "Results")
    matrix = results.pivot(index="model", columns="features", values="pr_auc")
    matrix = matrix.reindex(index=MODELS, columns=FEATURES)
    chart_data = CategoryChartData()
    chart_data.categories = list(MODELS)
    for feature, label in zip(FEATURES, FEATURE_LABELS):
        chart_data.add_series(label, matrix[feature].astype(float).tolist())

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.75),
        Inches(1.62),
        Inches(11.8),
        Inches(4.65),
        chart_data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 1
    chart.value_axis.major_unit = 0.2
    chart.value_axis.tick_labels.number_format = "0.0"
    chart.category_axis.tick_labels.font.name = "Aptos"
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.value_axis.tick_labels.font.name = "Aptos"
    chart.value_axis.tick_labels.font.size = Pt(10)
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "PR-AUC"
    chart.plots[0].gap_width = 60
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    chart.plots[0].data_labels.number_format = "0.000"
    chart.plots[0].data_labels.font.size = Pt(9)
    for series, color in zip(chart.series, PALETTE):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color

    _add_text(
        slide,
        "The validation/test partitions remain full size; all matrix models share the same 50,000-row training sample.",
        0.85,
        6.36,
        11.5,
        0.38,
        size=13,
        color=MUTED,
    )


def create_deck() -> Path:
    if not RESULTS_PATH.is_file():
        raise FileNotFoundError(
            f"Missing {RESULTS_PATH}. Run the experiment pipeline before creating the presentation."
        )
    if not METADATA_PATH.is_file():
        raise FileNotFoundError(
            f"Missing {METADATA_PATH}. Run scripts/run_sampled_graph_features.py first."
        )

    results = pd.read_csv(RESULTS_PATH)
    _validate_results(results)
    metadata = json.loads(METADATA_PATH.read_text(encoding="utf-8"))
    sample_edges = int(metadata["sampled_training_edges"])
    sample_frauds = int(metadata["sampled_training_frauds"])
    sample_accounts = int(metadata["sampled_training_accounts"])
    seed = int(metadata["seed"])

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, NAVY)
    for x, color, size in ((9.5, BLUE, 2.0), (10.45, TEAL, 1.45), (11.2, GOLD, 0.95)):
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(1.2 + (x - 9.5) * 0.35),
            Inches(size), Inches(size)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
    _add_text(slide, "GRAPH-BASED", 0.8, 1.25, 7.7, 0.65, size=20, color=TEAL, bold=True)
    _add_text(slide, "Fraud Detection", 0.8, 2.0, 8.5, 1.0, size=40, color=WHITE, bold=True)
    _add_text(
        slide,
        "A resource-bounded local graph experiment on PaySim",
        0.84,
        3.18,
        7.4,
        0.8,
        size=21,
        color=RGBColor(207, 220, 236),
    )
    _add_text(slide, "12 matched experiments  |  4 model families  |  3 feature sets",
              0.84, 5.85, 10.8, 0.4, size=15, color=WHITE)
    _add_footer(slide, 1, dark=True)

    # 2. Objective and implementation choice
    slide = _new_slide(prs, "Question and implementation choice", 2, "Project framing")
    _add_text(
        slide,
        "Do account relationships and recent transaction activity improve fraud detection beyond transaction-level predictors?",
        0.8, 1.55, 11.7, 0.95, size=24, color=NAVY, bold=True,
    )
    _add_card(slide, "Local by design", "Build an account graph in NetworkX from a deterministic training-edge sample. No graph database service is required.", 0.8, 3.0, 3.7, 2.25, BLUE)
    _add_card(slide, "Matched comparison", "Compare tabular-only, tabular + graph, and tabular + graph + 24-hour history.", 4.8, 3.0, 3.7, 2.25, TEAL)
    _add_card(slide, "Time-aware evaluation", "Preserve chronological validation and test periods; keep threshold selection on validation.", 8.8, 3.0, 3.7, 2.25, GOLD)

    # 3. Dataset and split
    slide = _new_slide(prs, "Data: severe imbalance, chronological split", 3, "Benchmark")
    _add_card(slide, "6.36M", "transactions in the synthetic PaySim dataset", 0.8, 1.7, 3.7, 1.55, BLUE, 15)
    _add_card(slide, "8,213", "fraud transactions (0.1291%)", 4.8, 1.7, 3.7, 1.55, TEAL, 15)
    _add_card(slide, "50,000", f"sampled train edges; {sample_frauds} frauds; {sample_accounts:,} accounts", 8.8, 1.7, 3.7, 1.55, GOLD, 15)
    _add_text(slide, "Chronological partitions by unique transaction step", 0.85, 3.65, 11.6, 0.4,
              size=19, color=NAVY, bold=True)
    _add_card(slide, "TRAIN  •  through step 520", "6,082,007 rows total\n50,000 rows used in matched runs", 0.8, 4.25, 3.7, 1.5, BLUE, 15)
    _add_card(slide, "VALIDATION  •  steps 521–631", "191,147 rows\nFull period; used for early stopping and threshold selection", 4.8, 4.25, 3.7, 1.5, TEAL, 14)
    _add_card(slide, "TEST  •  steps 632–743", "89,466 rows\nFull period; held out for reported metrics", 8.8, 4.25, 3.7, 1.5, GOLD, 14)
    _add_text(slide, f"Sample seed: {seed}  •  sampled training data is about 0.82% of the full training split",
              0.85, 6.1, 11.5, 0.35, size=13, color=MUTED)

    # 4. Pipeline
    slide = _new_slide(prs, "End-to-end experimental workflow", 4, "Method")
    steps = [
        ("01", "Chronological split", "Train / validation / test by time"),
        ("02", "Feature building", "Tabular + causal temporal windows"),
        ("03", "Local graph sample", f"{sample_edges:,} training edges; seed {seed}"),
        ("04", "Train & select", "Four classifiers; validation threshold"),
        ("05", "Test & report", "PR-AUC and threshold-based metrics"),
    ]
    for index, (number, title, body) in enumerate(steps):
        x = 0.65 + index * 2.55
        _add_card(slide, f"{number}  {title}", body, x, 2.25, 2.25, 2.25, PALETTE[index % 3], 15)
        if index < len(steps) - 1:
            _add_text(slide, "→", x + 2.25, 3.1, 0.3, 0.45, size=21, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, "Scripts: preprocessing → temporal features → sampled graph features → experiments",
              0.85, 5.35, 11.5, 0.55, size=17, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    # 5. Graph features and leakage controls
    slide = _new_slide(prs, "Graph features with temporal safeguards", 5, "Method")
    _add_card(slide, "Account network", "Directed edge: origin account → destination account\nTransaction amount is the edge weight\nNetworkX weighted PageRank + weakly connected components", 0.8, 1.75, 5.6, 2.05, BLUE, 16)
    _add_card(slide, "Per-account attributes", "In/out degree\nWeighted degree\nUnique counterparties\nUnseen-account indicator", 6.85, 1.75, 5.6, 2.05, TEAL, 16)
    _add_card(slide, "Leakage control", "Training aggregates use only earlier steps; each step is updated after its features are formed. Validation/test use the frozen training graph.", 0.8, 4.25, 5.6, 1.75, GOLD, 15)
    _add_card(slide, "Temporal windows", "Causal range: [step − window, step). Current-step activity is excluded. Primary window is 24 hours; 6h and 72h caches support sensitivity runs.", 6.85, 4.25, 5.6, 1.75, BLUE, 15)

    # 6. Matrix and scoring
    slide = _new_slide(prs, "Matched experiment matrix and scoring", 6, "Evaluation")
    _add_text(slide, "4 model families × 3 feature sets = 12 main experiments", 0.8, 1.55, 11.7, 0.5,
              size=24, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    _add_card(slide, "Models", "Logistic Regression\nRandom Forest\nXGBoost\nMLP", 0.8, 2.4, 3.7, 2.35, BLUE, 16)
    _add_card(slide, "Feature sets", "Sampled tabular\nSampled tabular + graph\nSampled tabular + graph + 24h", 4.8, 2.4, 3.7, 2.35, TEAL, 15)
    _add_card(slide, "Evaluation", "Primary: test PR-AUC\nAlso: ROC-AUC, precision, recall, F1\nThreshold selected on validation F1", 8.8, 2.4, 3.7, 2.35, GOLD, 15)
    _add_text(slide, "Every matrix cell uses the same 50,000 sampled training row IDs; validation and test remain full-sized.",
              1.0, 5.45, 11.3, 0.65, size=17, color=INK, align=PP_ALIGN.CENTER)

    # 7. Result chart
    _add_chart_slide(prs, results, 7)

    # 8. Findings
    matrix = results.pivot(index="model", columns="features", values="pr_auc")
    graph_delta = matrix["sampled_graph"] - matrix["tabular_sampled"]
    temporal_delta = matrix["sampled_temporal"] - matrix["sampled_graph"]
    total_delta = matrix["sampled_temporal"] - matrix["tabular_sampled"]
    best_model = matrix["sampled_temporal"].idxmax()
    slide = _new_slide(prs, "What the measured results say", 8, "Findings")
    _add_card(slide, "Best score", f"{best_model}: {matrix.loc[best_model, 'sampled_temporal']:.6f} PR-AUC with graph + 24h features.", 0.8, 1.75, 3.7, 1.75, TEAL, 16)
    _add_card(slide, "No universal graph lift", f"Graph-only PR-AUC improved for {int((graph_delta > 0).sum())}/4 models; mean change {graph_delta.mean():+.6f}.", 4.8, 1.75, 3.7, 1.75, BLUE, 16)
    _add_card(slide, "Model-specific effects", f"Graph + temporal vs sampled tabular: MLP {total_delta['MLP']:+.6f}; Logistic Regression {total_delta['Logistic']:+.6f}.", 8.8, 1.75, 3.7, 1.75, GOLD, 15)
    _add_bullets(
        slide,
        [
            f"Adding 24-hour activity to graph features improved {int((temporal_delta > 0).sum())}/4 model families; mean change {temporal_delta.mean():+.6f}.",
            "Random Forest is consistently near the top on this synthetic benchmark; an almost-perfect score is not a real-world performance guarantee.",
            "Compare feature sets within a model and sample design; supplemental full-training tabular results are not a matched baseline.",
        ],
        1.0, 4.15, 11.3, 2.0, size=16, spacing=11,
    )

    # 9. Local implementation
    slide = _new_slide(prs, "A graph workflow that fits local resources", 9, "Implementation")
    _add_card(slide, "No database service", "Python + NetworkX; no Neo4j, AuraDB, Java service, or GDS installation in the experiment path.", 0.8, 1.8, 3.7, 2.0, BLUE, 16)
    _add_card(slide, "Bounded training graph", f"Deterministic cap: {sample_edges:,} edges\n{sample_accounts:,} accounts\nSeed: {seed}", 4.8, 1.8, 3.7, 2.0, TEAL, 16)
    _add_card(slide, "Observed run profile", "Prior local graph build used roughly 2.5–2.7 GB process memory on the project machine. This is an observed run, not a formal benchmark.", 8.8, 1.8, 3.7, 2.0, GOLD, 14)
    _add_text(slide, "Trade-off", 0.9, 4.55, 1.8, 0.35, size=17, color=TEAL, bold=True)
    _add_text(slide, "Lower setup and resource burden, but the sample cannot represent every edge in the full transaction graph.",
              0.9, 4.98, 11.4, 0.75, size=22, color=NAVY, bold=True)

    # 10. Limitations and next steps
    slide = _new_slide(prs, "Limitations and next experiments", 10, "Interpretation")
    _add_card(slide, "Benchmark scope", "Synthetic PaySim data; only 50,000 of 6.08M training rows enter the matched runs; one fixed random seed.", 0.8, 1.75, 5.6, 1.85, BLUE, 15)
    _add_card(slide, "Model scope", "Feature-engineered graph signals with conventional models. No graph neural network is implemented.", 6.85, 1.75, 5.6, 1.85, TEAL, 15)
    _add_bullets(
        slide,
        [
            "Repeat the matrix over multiple sample seeds and report variability.",
            "Evaluate on independent datasets or operationally realistic data before making deployment claims.",
            "Investigate the model-specific effect of graph and temporal features, including calibration and cost-sensitive thresholds.",
        ],
        1.0, 4.25, 11.2, 1.9, size=17, spacing=12,
    )

    # 11. Reproduction
    slide = _new_slide(prs, "Reproduce the project", 11, "Runbook")
    _add_text(slide, "Install and prepare", 0.9, 1.65, 3.0, 0.35, size=18, color=TEAL, bold=True)
    _add_text(slide, "python -m pip install -r requirements.txt\nPlace PaySim CSV at data.raw_path",
              0.9, 2.1, 11.3, 0.8, size=18, color=INK, font="Consolas")
    _add_text(slide, "Run from the repository root", 0.9, 3.2, 5.0, 0.35, size=18, color=TEAL, bold=True)
    commands = (
        "python scripts/run_preprocessing.py\n"
        "python scripts/run_temporal_features.py\n"
        "python scripts/run_sampled_graph_features.py\n"
        "python scripts/run_experiments.py\n"
        "python scripts/create_presentation.py"
    )
    command_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(3.7), Inches(11.6), Inches(2.15)
    )
    command_box.fill.solid()
    command_box.fill.fore_color.rgb = PALE
    command_box.line.fill.background()
    _add_text(slide, commands, 1.15, 3.95, 10.9, 1.7, size=16, color=INK, font="Consolas")
    _add_text(slide, "Raw data, caches, and generated metrics are intentionally excluded from Git.",
              0.9, 6.15, 11.4, 0.35, size=13, color=MUTED)

    # 12. Closing and source
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, NAVY)
    _add_text(slide, "TAKEAWAY", 0.8, 0.9, 3.0, 0.4, size=12, color=TEAL, bold=True)
    _add_text(
        slide,
        "Local graph features are feasible to evaluate—\nbut their value depends on the model and the sample.",
        0.8, 1.65, 11.5, 1.55, size=30, color=WHITE, bold=True,
    )
    _add_text(
        slide,
        "Measured results use a deterministic 50,000-edge training sample and full chronological validation/test periods. They are not full-graph or deployment estimates.",
        0.85, 4.05, 11.3, 0.9, size=17, color=RGBColor(207, 220, 236),
    )
    _add_text(
        slide,
        "Dataset: PaySim — Kaggle, Synthetic Financial Datasets For Fraud Detection (CC BY-SA 4.0)\nhttps://www.kaggle.com/datasets/ealaxi/paysim1",
        0.85, 5.55, 11.5, 0.7, size=12, color=RGBColor(183, 197, 214),
    )
    _add_footer(slide, 12, dark=True)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)
    return OUTPUT_PATH


def main() -> None:
    output = create_deck()
    print(
        f"Saved presentation: {output.relative_to(ROOT)} "
        f"({len(Presentation(output).slides)} slides)"
    )


if __name__ == "__main__":
    main()
